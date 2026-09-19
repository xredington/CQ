#!/usr/bin/env python3
"""Rebuild slides as native, fully editable PowerPoint shapes from a JSON spec.

    python3 rebuild.py deck.pptx spec.json --out deck_rebuilt.pptx [--theme t.json]

Every shape it emits is a real text box / rectangle / table — no images — so
the result is editable in PowerPoint like any hand-made slide.
See references/slide-spec.md for the spec format and the layout catalogue.
"""
import argparse, json, os, sys
from copy import deepcopy

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import theme as TH
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

IN = 914400


class Canvas:
    """Slide-sized drawing surface with the deck's own margins."""

    def __init__(self, slide, prs, t):
        self.s, self.t = slide, t
        self.W, self.H = prs.slide_width / IN, prs.slide_height / IN
        self.mx, self.mt, self.mb = t["margin_x"], t["margin_top"], t["margin_bot"]
        self.cw = self.W - 2 * self.mx          # content width
        self.y = self.mt                        # running cursor

    # -- primitives -----------------------------------------------------
    def tb(self, x, y, w, h):
        return TH.textbox(self.s.shapes, x, y, w, h)

    def rect(self, x, y, w, h, **kw):
        kw.setdefault("t", self.t)
        return TH.box(self.s.shapes, x, y, w, h, **kw)

    def line(self, x, y, w, color="hairline", h=0.012):
        return self.rect(x, y, w, h, fill=color)

    def text(self, frame, runs, **kw):
        kw.setdefault("t", self.t)
        return TH.set_text(frame, runs, **kw)

    # -- chrome ---------------------------------------------------------
    def header(self, spec):
        t = self.t
        y = self.mt
        if spec.get("eyebrow"):
            b = self.tb(self.mx, y, self.cw, 0.26)
            self.text(b.text_frame, spec["eyebrow"].upper(), size=t["size_eyebrow"],
                      color="primary", bold=True, font=t["font_head"])
            y += 0.3
        if spec.get("title"):
            h = 0.62 if len(spec["title"]) < 60 else 1.0
            b = self.tb(self.mx, y, self.cw, h)
            b.text_frame.word_wrap = True
            self.text(b.text_frame, spec["title"], size=t["size_title"],
                      color="ink", bold=True, font=t["font_head"], line=1.05)
            y += h + 0.04
        if spec.get("subtitle"):
            b = self.tb(self.mx, y, self.cw, 0.4)
            self.text(b.text_frame, spec["subtitle"], size=t["size_sub"],
                      color="ink_soft", line=1.2)
            y += 0.46
        self.rect(self.mx, y + 0.04, 1.15, 0.05, fill="primary")
        self.y = y + 0.34
        return self.y

    def footer(self, spec, number=None):
        t = self.t
        base = self.H - self.mb
        if spec.get("footnote"):
            b = self.tb(self.mx, base - 0.02, self.cw - 0.8, 0.3)
            self.text(b.text_frame, spec["footnote"], size=t["size_small"],
                      color="ink_soft")
        if number is not None:
            b = self.tb(self.W - self.mx - 0.7, base - 0.02, 0.7, 0.3)
            self.text(b.text_frame, str(number), size=t["size_small"],
                      color="ink_soft", align="right")

    @property
    def body_h(self):
        return self.H - self.mb - self.y - 0.12


# ----------------------------------------------------------------------
# layouts
# ----------------------------------------------------------------------
def _bullets(c, x, y, w, h, items, size=None, color="ink"):
    t = c.t
    size = size or t["size_body"]
    tb = c.tb(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        sub = isinstance(item, dict)
        txt = item.get("text", "") if sub else str(item)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "•  " + txt if not (sub and item.get("level")) else "–  " + txt
        p.space_after = Pt(9)
        p.line_spacing = 1.22
        if sub and item.get("level"):
            p.level = 1
        for r in p.runs:
            r.font.size = Pt(size - (2 if sub and item.get("level") else 0))
            r.font.name = t["font"]
            r.font.color.rgb = TH.rgb(t, "ink_soft" if (sub and item.get("level")) else color)
    return tb


def bullets_h(items, w, size, gap=9):
    """Height a bullet block needs, in inches."""
    h = 0.0
    for it in items:
        txt = it.get("text", "") if isinstance(it, dict) else str(it)
        sz = size - (2 if isinstance(it, dict) and it.get("level") else 0)
        h += TH.text_h("\u2022  " + txt, w, sz, line=1.22, pad=0) + gap / 72.0
    return h


def lay_bullets(c, spec):
    c.header(spec)
    items = spec.get("bullets", [])
    img = spec.get("side_note")
    w = c.cw * (0.6 if img else 1.0)
    _bullets(c, c.mx, c.y, w, c.body_h, items)
    if img:
        x = c.mx + c.cw * 0.64
        bw = c.cw * 0.36
        c.rect(x, c.y, bw, c.body_h, fill="primary_lt", radius=0.04)
        b = c.tb(x + 0.26, c.y + 0.26, bw - 0.52, c.body_h - 0.52)
        c.text(b.text_frame, img if isinstance(img, list) else [img],
               size=c.t["size_body"], color="primary_dk", line=1.25, space_after=8)


def lay_two_column(c, spec):
    c.header(spec)
    cols = spec.get("columns", [])[:2]
    gut = c.t["gutter"]
    w = (c.cw - gut) / 2
    for i, col in enumerate(cols):
        x = c.mx + i * (w + gut)
        y = c.y
        if col.get("heading"):
            b = c.tb(x, y, w, 0.38)
            c.text(b.text_frame, col["heading"], size=c.t["size_sub"],
                   bold=True, color="primary_dk", font=c.t["font_head"])
            c.line(x, y + 0.4, w * 0.35, "primary")
            y += 0.6
        _bullets(c, x, y, w, c.H - c.mb - y - 0.1, col.get("bullets", []))


def lay_cards(c, spec):
    c.header(spec)
    cards = spec.get("cards", [])[:4]
    if not cards:
        return
    t, gut = c.t, c.t["gutter"]
    n = len(cards)
    w = (c.cw - gut * (n - 1)) / n
    iw = w - 0.44                      # inner (padded) width
    hsize, bsize = t["size_sub"], t["size_body"] - 2

    # size every card to the tallest content, so the row reads as one band
    need = 0.0
    for card in cards:
        body = card.get("body", [])
        body = [body] if isinstance(body, str) else body
        blk = (bullets_h(card["bullets"], iw, bsize) if card.get("bullets")
               else sum(TH.text_h(x, iw, bsize, line=1.25) for x in body))
        need = max(need, 0.3 + (0.56 if card.get("badge") else 0)
                   + TH.text_h(card.get("heading", ""), iw, hsize, line=1.1, bold=True)
                   + 0.12 + blk + 0.26)
    h = min(c.body_h, max(need, 1.6))
    c.y += max(0.0, (c.body_h - h) / 2)

    for i, card in enumerate(cards):
        x = c.mx + i * (w + gut)
        c.rect(x, c.y, w, h, fill="surface_alt", line="hairline", radius=0.05)
        c.rect(x, c.y, w, 0.07, fill=TH.SERIES[i % len(TH.SERIES)])
        y = c.y + 0.3
        if card.get("badge"):
            bd = c.rect(x + 0.22, y, 0.42, 0.42, fill="primary", radius=0.5)
            c.text(bd.text_frame, str(card["badge"]), size=13, bold=True,
                   color="on_dark", align="center", font=t["font_head"])
            bd.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            y += 0.56
        hh = TH.text_h(card.get("heading", ""), iw, hsize, line=1.1, bold=True)
        b = c.tb(x + 0.22, y, iw, hh)
        c.text(b.text_frame, card.get("heading", ""), size=hsize, bold=True,
               color="ink", font=t["font_head"], line=1.1)
        y += hh + 0.12
        avail = c.y + h - y - 0.2
        if card.get("bullets"):
            _bullets(c, x + 0.22, y, iw, avail, card["bullets"], size=bsize)
        else:
            body = card.get("body", [])
            if body:
                b = c.tb(x + 0.22, y, iw, avail)
                c.text(b.text_frame, [body] if isinstance(body, str) else body,
                       size=bsize, color="ink_soft", line=1.25, space_after=6)


def lay_kpi(c, spec):
    c.header(spec)
    kpis = spec.get("kpis", [])[:5]
    if not kpis:
        return
    if not spec.get("bullets"):
        c.y += max(0.0, (c.body_h - 1.72) / 2)
    gut = c.t["gutter"]
    n = len(kpis)
    w = (c.cw - gut * (n - 1)) / n
    h = 1.72
    for i, k in enumerate(kpis):
        x = c.mx + i * (w + gut)
        c.rect(x, c.y, w, h, fill="primary_lt", radius=0.06)
        b = c.tb(x + 0.16, c.y + 0.22, w - 0.32, 0.72)
        c.text(b.text_frame, k.get("value", ""), size=c.t["size_kpi"], bold=True,
               color="primary_dk", align="center", font=c.t["font_head"])
        b = c.tb(x + 0.16, c.y + 0.98, w - 0.32, 0.62)
        c.text(b.text_frame, k.get("label", ""), size=c.t["size_body"] - 2,
               color="ink_soft", align="center", line=1.15)
    rest = c.y + h + 0.3
    if spec.get("bullets"):
        _bullets(c, c.mx, rest, c.cw, c.H - c.mb - rest - 0.1, spec["bullets"])


def lay_timeline(c, spec):
    """Phased roadmap — the workhorse for 'from summit to execution' slides."""
    c.header(spec)
    steps = spec.get("steps", [])[:6]
    if not steps:
        return
    t = c.t
    n, gut = len(steps), 0.18
    w = (c.cw - gut * (n - 1)) / n
    iw = w - 0.36
    hsize, bsize = t["size_body"] + 1, t["size_body"] - 3
    rail_y = c.y + 0.46
    c.rect(c.mx, rail_y, c.cw, 0.035, fill="hairline")
    top = rail_y + 0.34

    need = 0.0
    for s_ in steps:
        blk = (bullets_h(s_["bullets"], iw, bsize, gap=7) if s_.get("bullets")
               else TH.text_h(s_.get("body", ""), iw, bsize, line=1.2))
        need = max(need, 0.22 + TH.text_h(s_.get("heading", ""), iw, hsize,
                                          line=1.1, bold=True) + 0.1
                   + (0.32 if s_.get("date") else 0) + blk + 0.24)
    h = min(c.H - c.mb - top - 0.1, max(need, 1.4))

    for i, s_ in enumerate(steps):
        x = c.mx + i * (w + gut)
        col = TH.SERIES[i % len(TH.SERIES)]
        if s_.get("phase"):
            b = c.tb(x, c.y, w, 0.34)
            c.text(b.text_frame, str(s_["phase"]).upper(), size=t["size_eyebrow"],
                   bold=True, color=col, font=t["font_head"])
        c.rect(x, rail_y, w * 0.98, 0.035, fill=col)
        c.rect(x, rail_y - 0.1, 0.24, 0.24, fill=col, radius=0.5)
        c.rect(x, top, w, h, fill="surface_alt", line="hairline", radius=0.05)
        y = top + 0.22
        hh = TH.text_h(s_.get("heading", ""), iw, hsize, line=1.1, bold=True)
        b = c.tb(x + 0.18, y, iw, hh)
        c.text(b.text_frame, s_.get("heading", ""), size=hsize, bold=True,
               color="ink", font=t["font_head"], line=1.1)
        y += hh + 0.1
        if s_.get("date"):
            b = c.tb(x + 0.18, y, iw, 0.28)
            c.text(b.text_frame, s_["date"], size=t["size_small"], color=col, bold=True)
            y += 0.32
        avail = top + h - y - 0.16
        if s_.get("bullets"):
            _bullets(c, x + 0.18, y, iw, avail, s_["bullets"], size=bsize)
        elif s_.get("body"):
            b = c.tb(x + 0.18, y, iw, avail)
            c.text(b.text_frame, s_["body"], size=bsize, color="ink_soft", line=1.2)


def lay_process(c, spec):
    c.header(spec)
    steps = spec.get("steps", [])[:6]
    n = len(steps)
    if not n:
        return
    gut = 0.1
    w = (c.cw - gut * (n - 1)) / n
    h = 1.05
    for i, s in enumerate(steps):
        x = c.mx + i * (w + gut)
        shp = c.s.shapes.add_shape(
            MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON,
            Emu(int(x * IN)), Emu(int(c.y * IN)),
            Emu(int(w * IN)), Emu(int(h * IN)))
        shp.fill.solid()
        shp.fill.fore_color.rgb = TH.rgb(c.t, "primary" if i % 2 == 0 else "primary_dk")
        shp.line.fill.background()
        shp.shadow.inherit = False
        TH.no_shadow(shp)
        tf = shp.text_frame
        tf.word_wrap = True
        c.text(tf, s.get("heading", s if isinstance(s, str) else ""),
               size=c.t["size_body"] - 1, bold=True, color="on_dark", align="center")
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    y = c.y + h + 0.3
    descs = [s.get("body", "") for s in steps if isinstance(s, dict)]
    if any(descs):
        for i, d in enumerate(descs):
            x = c.mx + i * (w + gut)
            b = c.tb(x, y, w, c.H - c.mb - y - 0.1)
            c.text(b.text_frame, d, size=c.t["size_body"] - 3, color="ink_soft",
                   align="center", line=1.2)


def lay_matrix(c, spec):
    c.header(spec)
    q = spec.get("quadrants", [])[:4]
    gut = c.t["gutter"]
    w = (c.cw - gut) / 2
    h = (min(c.body_h, 4.4) - gut) / 2
    for i, item in enumerate(q):
        x = c.mx + (i % 2) * (w + gut)
        y = c.y + (i // 2) * (h + gut)
        col = TH.SERIES[i % len(TH.SERIES)]
        c.rect(x, y, w, h, fill="surface_alt", line="hairline", radius=0.04)
        c.rect(x, y, 0.07, h, fill=col)
        b = c.tb(x + 0.26, y + 0.2, w - 0.5, 0.42)
        c.text(b.text_frame, item.get("heading", ""), size=c.t["size_sub"],
               bold=True, color="ink", font=c.t["font_head"])
        if item.get("bullets"):
            _bullets(c, x + 0.26, y + 0.66, w - 0.5, h - 0.86, item["bullets"],
                     size=c.t["size_body"] - 2)
        elif item.get("body"):
            b = c.tb(x + 0.26, y + 0.66, w - 0.5, h - 0.86)
            c.text(b.text_frame, item["body"], size=c.t["size_body"] - 2,
                   color="ink_soft", line=1.22)


def lay_table(c, spec):
    c.header(spec)
    rows = spec.get("rows", [])
    head = spec.get("header", [])
    if not rows:
        return
    ncol = len(head or rows[0])
    nrow = len(rows) + (1 if head else 0)
    h = min(c.body_h, 0.42 * nrow + 0.2)
    gfx = c.s.shapes.add_table(nrow, ncol, Emu(int(c.mx * IN)), Emu(int(c.y * IN)),
                               Emu(int(c.cw * IN)), Emu(int(h * IN)))
    tbl = gfx.table
    tbl.first_row = bool(head)
    data = ([head] if head else []) + rows
    for r, row in enumerate(data):
        for col in range(ncol):
            cell = tbl.cell(r, col)
            cell.text = str(row[col]) if col < len(row) else ""
            cell.margin_left = Emu(int(0.1 * IN))
            cell.margin_top = Emu(int(0.05 * IN))
            is_head = head and r == 0
            cell.fill.solid()
            cell.fill.fore_color.rgb = TH.rgb(
                c.t, "primary" if is_head else ("surface" if r % 2 else "surface_alt"))
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(c.t["size_body"] - 3)
                    run.font.name = c.t["font"]
                    run.font.bold = bool(is_head)
                    run.font.color.rgb = TH.rgb(c.t, "on_dark" if is_head else "ink")


def lay_section(c, spec):
    t = c.t
    c.rect(0, 0, c.W, c.H, fill="primary_dk")
    c.rect(0, c.H - 0.28, c.W, 0.28, fill="primary")
    y = c.H / 2 - 0.9
    if spec.get("eyebrow"):
        b = c.tb(c.mx + 0.3, y, c.cw, 0.32)
        c.text(b.text_frame, spec["eyebrow"].upper(), size=t["size_eyebrow"] + 1,
               bold=True, color="primary_lt", font=t["font_head"])
        y += 0.42
    b = c.tb(c.mx + 0.3, y, c.cw - 0.6, 1.2)
    c.text(b.text_frame, spec.get("title", ""), size=t["size_title"] + 10,
           bold=True, color="on_dark", font=t["font_head"], line=1.05)
    if spec.get("subtitle"):
        b = c.tb(c.mx + 0.3, y + 1.3, c.cw - 0.6, 0.6)
        c.text(b.text_frame, spec["subtitle"], size=t["size_sub"],
               color="primary_lt", line=1.25)


def lay_quote(c, spec):
    t = c.t
    c.rect(0, 0, c.W, c.H, fill="surface_alt")
    c.rect(c.mx, c.H / 2 - 1.3, 0.08, 2.6, fill="primary")
    b = c.tb(c.mx + 0.42, c.H / 2 - 1.25, c.cw - 0.8, 1.8)
    c.text(b.text_frame, spec.get("quote", spec.get("title", "")),
           size=t["size_title"] - 4, color="ink", font=t["font_head"], line=1.22)
    if spec.get("attribution"):
        b = c.tb(c.mx + 0.42, c.H / 2 + 0.72, c.cw - 0.8, 0.5)
        c.text(b.text_frame, spec["attribution"], size=t["size_body"],
               color="primary_dk", bold=True)


LAYOUTS = {
    "bullets": lay_bullets, "title-bullets": lay_bullets,
    "two-column": lay_two_column, "columns": lay_two_column,
    "cards": lay_cards, "kpi": lay_kpi, "kpis": lay_kpi,
    "timeline": lay_timeline, "roadmap": lay_timeline,
    "process": lay_process, "matrix": lay_matrix,
    "table": lay_table, "section": lay_section, "quote": lay_quote,
}


# ----------------------------------------------------------------------
def clear_slide(slide, media_dir=None, idx=0):
    """Remove every shape, saving any pictures to media_dir first."""
    saved = []
    for shp in list(slide.shapes):
        if media_dir and shp.shape_type == 13:
            try:
                img = shp.image
                fn = os.path.join(media_dir, f"orig_slide{idx:03d}_{len(saved)+1}.{img.ext}")
                with open(fn, "wb") as fh:
                    fh.write(img.blob)
                saved.append(fn)
            except Exception:
                pass
        shp._element.getparent().remove(shp._element)
    return saved


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("deck")
    ap.add_argument("spec", help="JSON: {'slides':[{index, layout, ...}]}")
    ap.add_argument("--out", required=True)
    ap.add_argument("--theme")
    ap.add_argument("--keep-media", help="dir to save replaced images into")
    ap.add_argument("--number", action="store_true", help="stamp slide numbers")
    a = ap.parse_args()

    t = TH.load(a.theme)
    with open(a.spec) as fh:
        spec = json.load(fh)
    if isinstance(spec, list):
        spec = {"slides": spec}
    if a.keep_media:
        os.makedirs(a.keep_media, exist_ok=True)

    prs = Presentation(a.deck)
    slides = list(prs.slides)
    done = []
    for s in spec["slides"]:
        i = int(s["index"])
        if not 1 <= i <= len(slides):
            print(f"!! slide {i} out of range (deck has {len(slides)}) — skipped")
            continue
        layout = s.get("layout", "bullets")
        if layout not in LAYOUTS:
            print(f"!! slide {i}: unknown layout '{layout}' — skipped")
            continue
        slide = slides[i - 1]
        clear_slide(slide, a.keep_media, i)
        c = Canvas(slide, prs, t)
        LAYOUTS[layout](c, s)
        if layout not in ("section", "quote"):
            c.footer(s, i if a.number else None)
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]
        done.append((i, layout))

    prs.save(a.out)
    for i, l in done:
        print(f"  slide {i:>3}  ->  {l}")
    print(f"rebuilt {len(done)} slide(s) -> {a.out}")


if __name__ == "__main__":
    main()
