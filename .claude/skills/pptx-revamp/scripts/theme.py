"""Brand tokens + shared drawing helpers for deck rebuilds.

Colours default to the CloudQuarks / Redington palette already used by the
prototypes in this repo. Override any of them with a JSON file passed as
--theme; keys mirror the TOKENS dict below.
"""
import json
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

TOKENS = {
    # --- palette -------------------------------------------------------
    "ink":        "04120A",  # near-black green, primary text
    "ink_soft":   "63706A",  # secondary text
    "primary":    "0A9D4E",  # Redington green
    "primary_dk": "065132",
    "primary_lt": "EAF4EE",
    "accent":     "F44609",  # orange, use sparingly for emphasis
    "info":       "0B6FD0",  # blue
    "info_dk":    "7C5CD6",  # violet
    # --- the categorical set: six vivid hues + their tints --------------
    "c1": "0B6FD0", "c1_lt": "E4F0FC",   # blue
    "c2": "0A9D4E", "c2_lt": "E3F6EB",   # green
    "c3": "F5A623", "c3_lt": "FEF3E0",   # amber
    "c4": "F44609", "c4_lt": "FEEAE3",   # coral
    "c5": "7C5CD6", "c5_lt": "EFEAFB",   # violet
    "c6": "12A4A4", "c6_lt": "E2F5F5",   # teal
    "surface":    "FFFFFF",
    "surface_alt":"F5F8F6",
    "hairline":   "D7DDD9",
    "on_dark":    "FFFFFF",
    # --- type ----------------------------------------------------------
    "font":       "Calibri",        # safe cross-platform default
    "font_head":  "Calibri",        # swap to a brand face if installed
    "size_title": 30,
    "size_sub":   16,
    "size_body":  15,
    "size_small": 11,
    "size_kpi":   36,
    "size_eyebrow": 11,
    # --- geometry (inches) ---------------------------------------------
    "margin_x":   0.62,
    "margin_top": 0.55,
    "margin_bot": 0.52,
    "gutter":     0.26,
}

# categorical series colours, in order, with the matching card tints
SERIES = ["c1", "c2", "c3", "c4", "c5", "c6"]
TINTS = ["c1_lt", "c2_lt", "c3_lt", "c4_lt", "c5_lt", "c6_lt"]


def series(t, i):
    return SERIES[i % len(SERIES)]


def tint(t, i):
    return TINTS[i % len(TINTS)]


def load(path=None):
    t = dict(TOKENS)
    if path:
        with open(path) as fh:
            t.update(json.load(fh))
    return t


def no_shadow(shape):
    """Strip the theme's default drop shadow — flat shapes render the same
    in PowerPoint, LibreOffice and Google Slides."""
    from pptx.oxml.ns import qn
    el = shape._element
    spPr = el.spPr
    for tag in ("a:effectLst", "a:effectRef"):
        for node in spPr.findall(qn(tag)):
            spPr.remove(node)
    spPr.append(spPr.makeelement(qn("a:effectLst"), {}))
    # the shape's style reference can re-apply a themed shadow; point it at
    # the theme's "no effect" slot instead
    style = el.find(qn("p:style"))
    if style is not None:
        ref = style.find(qn("a:effectRef"))
        if ref is not None:
            ref.set("idx", "0")
    return shape


def est_lines(text, width_in, pt, bold=False):
    """Rough line count for a string in a box of the given width.

    Deliberately pessimistic: over-reserving costs white space, while
    under-reserving makes a heading collide with the text under it. Bold and
    fallback faces (when Calibri is absent) run wider, hence the factors.
    """
    if not text:
        return 0
    per_line = max(6, int(width_in * 96 / (pt * (0.85 if bold else 0.72))))
    n = 0
    for para in str(text).split("\n"):
        n += max(1, -(-len(para) // per_line))
    return n


def fit_one_line(text, width_in, start_pt, min_pt=14, bold=True):
    """Largest point size at which `text` still fits on a single line.

    Measures width directly: est_lines() floors at six characters per line,
    which lies about short strings like "+22.9%" in a narrow column.
    """
    if not text:
        return start_pt
    # display numerals are near-tabular and run wide; measured against a
    # rendered deck, not guessed
    factor = 1.0 if bold else 0.8
    pt = start_pt
    while pt > min_pt and len(str(text)) * pt * factor / 96.0 > width_in:
        pt -= 1
    return pt


def text_h(text, width_in, pt, line=1.2, pad=0.08, bold=False):
    """Height in inches needed to render `text` at `pt` in `width_in`."""
    # 1.25 compensates for faces with a taller line box than Calibri's, so a
    # block measured here still fits when the deck opens elsewhere.
    return est_lines(text, width_in, pt, bold) * (pt * line * 1.25 / 72.0) + pad


def rgb(t, key):
    """Token name or literal hex -> RGBColor."""
    val = t.get(key, key)
    return RGBColor.from_string(str(val).lstrip("#").upper())


def set_text(frame, runs, t, size=None, color="ink", bold=False,
             align=None, font=None, line=None, space_after=0):
    """Fill a text frame. `runs` is a string or list of strings (one per para)."""
    from pptx.enum.text import PP_ALIGN
    if isinstance(runs, str):
        runs = [runs]
    frame.word_wrap = True
    for i, text in enumerate(runs):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = str(text)
        p.space_after = Pt(space_after)
        if line:
            p.line_spacing = line
        if align:
            p.alignment = getattr(PP_ALIGN, align.upper())
        for r in p.runs:
            r.font.size = Pt(size or t["size_body"])
            r.font.bold = bold
            r.font.name = font or t["font"]
            r.font.color.rgb = rgb(t, color)
    return frame


def box(shapes, x, y, w, h, fill=None, t=None, line=None, line_w=1,
        radius=None, shadow=False):
    """Rounded/plain rectangle with no default outline or shadow."""
    from pptx.enum.shapes import MSO_SHAPE
    shp = shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Emu(int(x * 914400)), Emu(int(y * 914400)),
        Emu(int(w * 914400)), Emu(int(h * 914400)))
    if radius is not None:
        try:
            shp.adjustments[0] = radius
        except (IndexError, ValueError):
            pass
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(t, fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(t, line)
        shp.line.width = Pt(line_w)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    if not shadow:
        no_shadow(shp)
    shp.text_frame.word_wrap = True
    return shp


def textbox(shapes, x, y, w, h):
    from pptx.util import Emu as E
    tb = shapes.add_textbox(E(int(x * 914400)), E(int(y * 914400)),
                            E(int(w * 914400)), E(int(h * 914400)))
    tb.text_frame.word_wrap = True
    return tb
